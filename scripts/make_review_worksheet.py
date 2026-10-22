"""Build data/universe/review_worksheet.csv/.xlsx for Caleb's manual review.
Input: titles_v2_all.csv (REVIEW rows + AUTO_KEEP rows with a judgment sector
override), sector_overrides.csv. Date-check rows come from a cik list (no Tiingo
values are read into or written to the worksheet: no prices, returns, or Tiingo dates).
Primary row_type = first matching type in PRIORITY; extra tasks in `tasks`."""
import pandas as pd
from urllib.parse import quote
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

DATE_CHECK = {"1884072", "1892274", "1934245", "2002236", "1993097"}   # Adamas One, Visionary, Newton Golf, Springview, 3 E Network
PRIORITY = [   # (row_type, reason, minutes, session)
    ("DATE_CHECK", None, 4, 1), ("SECTOR_ONLY", None, 1, 1),
    ("FOUND_CANDIDATE", "SEARCH_CANDIDATE", 2.5, 2), ("SEARCH_MISS_UNCERTAIN", "SEARCH_MISS", 3, 2),
    ("NAME_MISMATCH", "NAME_MISMATCH", 2, 3), ("LIST_PAGE", "LIST_PAGE", 2, 3), ("DISAMBIG", "DISAMBIG", 1.5, 3),
    ("SEARCH_UNVERIFIED", "SEARCH_UNVERIFIED", 1.5, 4),
    ("NO_PRE_IPO_DATA", "NO_PRE_IPO_DATA", 1.5, 5), ("NO_VIEWS", "NO_VIEWS", 1.5, 5),
    ("MANY_REDIRECTS", "MANY_REDIRECTS", 2, 5), ("FOREIGN_ADR", "FOREIGN_ADR", 1, 5),
]
SECTORS = ["Consumer Discretionary", "Consumer Staples", "Energy", "Financials", "Health Care", "Industrials",
           "Information Technology", "Materials", "Other", "Real Estate"]
ALLOWED = {
    "decision": ["keep", "drop", "unsure"],
    "sector_confirm": ["ok"] + SECTORS,
    "date_confirm": ["ok", "wrong", "unresolved"],
    "review_code": ["MATCH_FOUND_MANUAL", "CANDIDATE_CONFIRMED", "CANDIDATE_REJECTED", "NO_ARTICLE",
                    "SECTION_REDIRECT_ONLY", "ARTICLE_AFTER_IPO", "REDIRECT_OK", "WRONG_TARGET", "PARENT_SUB_FIXED",
                    "PARENT_SUB_PROXY", "OLD_TITLES_ADDED", "RENAME_IRRELEVANT", "LOW_VIEWS_REAL", "DISAMBIG_RESOLVED",
                    "FOREIGN_CHECKED", "DATE_CONFIRMED", "DATE_CORRECTED", "DATE_UNRESOLVED", "SECTOR_CONFIRMED",
                    "SECTOR_CHANGED", "NOT_REVIEWED_CAP"],
}
t = pd.read_csv("data/universe/titles_v2_all.csv", dtype=str).fillna("")
ov = pd.read_csv("data/universe/sector_overrides.csv", dtype=str).fillna("").set_index("cik")
rows = t[(t.decision == "REVIEW") | ((t.decision == "AUTO_KEEP") & t.cik.isin(ov.index))].copy()


def classify(r):
    rs = set(r.reasons.split())
    if r.cik in DATE_CHECK:
        return PRIORITY[0]
    if r.decision == "AUTO_KEEP":
        return PRIORITY[1]
    for p in PRIORITY[2:]:
        if p[1] in rs:
            return p
    return ("OTHER", None, 1.5, 5)


rows[["row_type", "_r", "est_minutes", "session"]] = rows.apply(lambda r: pd.Series(classify(r)), axis=1)
def tasks(r):
    out = ["TITLE"] if r.decision == "REVIEW" else []
    if r.cik in DATE_CHECK: out.append("DATE")
    if r.cik in ov.index: out.append("SECTOR")
    if "FOREIGN_ADR" in r.reasons.split(): out.append("ADR")
    return " ".join(out)
rows["tasks"] = rows.apply(tasks, axis=1)
rows["est_minutes"] = rows.est_minutes.astype(float) + rows.tasks.str.contains("SECTOR") * 0.5 * (rows.row_type != "SECTOR_ONLY") \
    + rows.tasks.str.contains("DATE") * 0 
rows["sector_override_reason"] = [ov.reason.get(c, "") for c in rows.cik]
q = lambda s: quote(s.replace(" ", "_"), safe="_()'")
rows["wikipedia_url"] = rows.final_title.apply(lambda s: f"https://en.wikipedia.org/wiki/{q(s)}" if s else "")
rows["candidate_url"] = rows.deep_best_title.apply(lambda s: f"https://en.wikipedia.org/wiki/{q(s)}" if s else "")
rows["wikipedia_search_url"] = rows.legal_name.apply(lambda s: "https://en.wikipedia.org/w/index.php?search=" + quote(s) + "&fulltext=1")
rows["wikidata_url"] = rows.wikidata_qid.apply(lambda s: f"https://www.wikidata.org/wiki/{s}" if s else "")
rows["wikidata_search_url"] = rows.legal_name.apply(lambda s: "https://www.wikidata.org/w/index.php?search=" + quote(s))
rows["edgar_company_url"] = rows.cik.apply(lambda c: f"https://www.sec.gov/edgar/browse/?CIK={int(c)}")
rows["edgar_filing_url"] = rows.filing_url
rows["edgar_8k_url"] = rows.cik.apply(lambda c: f"https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK={int(c)}&type=8-K&dateb=&owner=include&count=40")
rows["date_to_check"] = [r.first_trade_date if r.cik in DATE_CHECK else "" for r in rows.itertuples()]
for c in ("decision_review", "chosen_title", "sector_confirm", "date_confirm", "review_code_review", "notes"):
    rows[c] = ""
rows = rows.sort_values(["session", "est_minutes", "row_type", "year", "legal_name"]).reset_index(drop=True)
rows.insert(0, "order", range(1, len(rows) + 1))
out = rows.drop(columns=["review_code"]).rename(columns={"decision": "script_decision", "reasons": "script_reasons"})
out = out.rename(columns={"decision_review": "decision", "review_code_review": "review_code"})
cols = ["order", "session", "row_type", "tasks", "est_minutes", "year", "cik", "legal_name", "tickers", "exchanges",
        "form_type", "filing_date_424b4", "date_to_check", "first_trade_conf", "sector", "sector_override_reason",
        "sic", "sic_description", "foreign_flag", "attention_only", "attention_reason", "script_decision",
        "script_reasons", "final_title", "deep_class", "deep_best_title", "deep_candidates",
        "wikipedia_url", "candidate_url", "wikipedia_search_url", "wikidata_url", "wikidata_search_url",
        "edgar_company_url", "edgar_filing_url", "edgar_8k_url",
        "decision", "chosen_title", "sector_confirm", "date_confirm", "review_code", "notes"]
out = out[cols]
BANNED = ["price", "close", "open", "pop", "ret30", "return", "tiingo", "adj", "volume"]
bad = [c for c in out.columns if any(b in c.lower() for b in BANNED)]
assert not bad, bad
out.to_csv("data/universe/review_worksheet.csv", index=False)

wb = Workbook(); ws = wb.active; ws.title = "review"
ws.append(cols)
for r in out.itertuples(index=False):
    ws.append(list(r))
for c in ws[1]:
    c.font = Font(bold=True)
ws.freeze_panes = "I2"
n = len(out) + 1
lists = wb.create_sheet("allowed_values")
for j, (col, vals) in enumerate(ALLOWED.items(), 1):
    lists.cell(1, j, col).font = Font(bold=True)
    for i, v in enumerate(vals, 2):
        lists.cell(i, j, v)
    L = get_column_letter(j)
    dv = DataValidation(type="list", formula1=f"=allowed_values!${L}$2:${L}${len(vals) + 1}", allow_blank=True,
                        showErrorMessage=True, errorTitle="Not allowed", error=f"Pick a value from the {col} list")
    ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(cols.index(col) + 1)}2:{get_column_letter(cols.index(col) + 1)}{n}")
for col in [c for c in cols if c.endswith("_url")]:
    j = cols.index(col) + 1
    for i in range(2, n + 1):
        cell = ws.cell(i, j)
        if cell.value:
            cell.hyperlink = cell.value; cell.style = "Hyperlink"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n}"
wb.save("data/universe/review_worksheet.xlsx")
s = out.groupby(["session", "row_type"]).agg(rows=("cik", "size"), minutes=("est_minutes", "sum")).reset_index()
s.to_csv("verify/review_worksheet_summary.csv", index=False)
print(s.to_string()); print("total rows", len(out), "hours", round(out.est_minutes.sum() / 60, 1))
print("tasks:", out.tasks.str.split().explode().value_counts().to_dict())
print("banned-column check passed:", not bad)
